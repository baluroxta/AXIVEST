"""
AXIVEST — Portfolio Risk & Options Intelligence
Financial Risk Model & Executive Workbook Generator
Author: Kandala Guru Balaji
Degree: B.Tech Computer Science and Business Systems (3rd Year)
Tagline: "See the risk. Test the downside."

Builds: AXIVEST/excel/axivest_financial_risk_model.xlsx
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, "excel", "axivest_financial_risk_model.xlsx")

def create_financial_workbook():
    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active
    wb.remove(default_sheet)

    # Styles
    font_title = Font(name="Segoe UI", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="E2E8F0")
    font_section = Font(name="Segoe UI", size=13, bold=True, color="1B365D")
    font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Segoe UI", size=10, bold=True, color="1A202C")
    font_regular = Font(name="Segoe UI", size=10, color="2D3748")
    font_kpi_num = Font(name="Segoe UI", size=16, bold=True, color="1B365D")
    font_kpi_label = Font(name="Segoe UI", size=9, bold=True, color="718096")

    fill_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    fill_slate = PatternFill(start_color="2B4C7E", end_color="2B4C7E", fill_type="solid")
    fill_subtle_blue = PatternFill(start_color="EDF2F7", end_color="EDF2F7", fill_type="solid")
    fill_accent_kpi = PatternFill(start_color="F7FAFC", end_color="F7FAFC", fill_type="solid")
    fill_stripe = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_green_subtle = PatternFill(start_color="E6FFFA", end_color="E6FFFA", fill_type="solid")
    fill_red_subtle = PatternFill(start_color="FFF5F5", end_color="FFF5F5", fill_type="solid")

    thin_border_side = Side(style="thin", color="CBD5E0")
    border_all = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_kpi = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    def auto_fit_columns(ws, min_col=1, max_col=15):
        ws.views.sheetView[0].showGridLines = True
        for col in range(min_col, max_col + 1):
            col_letter = get_column_letter(col)
            max_len = 0
            for row in range(1, ws.max_row + 1):
                val = ws.cell(row=row, column=col).value
                if val is not None:
                    s_val = str(val)
                    if not s_val.startswith("="):
                        max_len = max(max_len, len(s_val))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # =========================================================================
    # TAB 1: EXECUTIVE SUMMARY & KPIS
    # =========================================================================
    ws1 = wb.create_sheet(title="Executive Summary")
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:G1")
    ws1["A1"] = "AXIVEST — PORTFOLIO RISK & OPTIONS INTELLIGENCE"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_navy
    ws1["A1"].alignment = align_center

    ws1.merge_cells("A2:G2")
    ws1["A2"] = "Author: Kandala Guru Balaji | B.Tech CSBS (3rd Year) | Tagline: \"See the risk. Test the downside.\""
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = fill_slate
    ws1["A2"].alignment = align_center
    ws1.row_dimensions[1].height = 30
    ws1.row_dimensions[2].height = 20

    # KPI Cards Section
    kpis = [
        ("Avg Annual Volatility", "=AVERAGE('Risk Rankings'!C6:C15)", "0.0%", "A", 4, 5),
        ("Worst Max Drawdown", "=MIN('Risk Rankings'!D6:D16)", "0.0%", "B", 4, 5),
        ("Assets Analysed", "11", "0", "C", 4, 5),
        ("Severe Stress Loss", "=ABS('Macro Stress Testing'!B23)", "0.0%", "D", 4, 5),
        ("Options Strategies", "3", "0", "E", 4, 5),
    ]

    for label, val, num_fmt, col_letter, r_top, r_bot in kpis:
        ws1[f"{col_letter}{r_top}"] = label
        ws1[f"{col_letter}{r_top}"].font = font_kpi_label
        ws1[f"{col_letter}{r_top}"].fill = fill_accent_kpi
        ws1[f"{col_letter}{r_top}"].alignment = align_center
        ws1[f"{col_letter}{r_top}"].border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side)
        
        ws1[f"{col_letter}{r_bot}"] = val
        ws1[f"{col_letter}{r_bot}"].font = font_kpi_num
        ws1[f"{col_letter}{r_bot}"].fill = fill_accent_kpi
        ws1[f"{col_letter}{r_bot}"].alignment = align_center
        ws1[f"{col_letter}{r_bot}"].number_format = num_fmt
        ws1[f"{col_letter}{r_bot}"].border = Border(left=thin_border_side, right=thin_border_side, bottom=thin_border_side)

    ws1.row_dimensions[4].height = 18
    ws1.row_dimensions[5].height = 26

    # Portfolio Overview Table
    ws1["A7"] = "Institutional Equity Universe & Benchmark Baseline"
    ws1["A7"].font = font_section

    headers1 = ["Asset Name", "Ticker", "Sector", "Annualized Volatility", "Max Drawdown", "Risk Category"]
    for col_idx, h in enumerate(headers1, 1):
        c = ws1.cell(row=8, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = border_all
    ws1.row_dimensions[8].height = 24

    asset_meta = [
        ("Bharti Airtel", "BHARTIARTL.NS", "Telecom", "='Risk Rankings'!C11", "='Risk Rankings'!D11", "Moderate Risk"),
        ("HDFC Bank", "HDFCBANK.NS", "Banking", "='Risk Rankings'!C13", "='Risk Rankings'!D13", "Core Large-Cap"),
        ("ICICI Bank", "ICICIBANK.NS", "Banking", "='Risk Rankings'!C15", "='Risk Rankings'!D15", "Core Large-Cap"),
        ("Infosys", "INFY.NS", "Information Tech", "='Risk Rankings'!C6", "='Risk Rankings'!D6", "High Volatility"),
        ("ITC", "ITC.NS", "FMCG", "='Risk Rankings'!C14", "='Risk Rankings'!D14", "Defensive Value"),
        ("L&T", "LT.NS", "Infrastructure", "='Risk Rankings'!C9", "='Risk Rankings'!D9", "Industrial Growth"),
        ("M&M", "M&M.NS", "Automobile", "='Risk Rankings'!C7", "='Risk Rankings'!D7", "High Volatility"),
        ("NTPC", "NTPC.NS", "Energy & Utilities", "='Risk Rankings'!C10", "='Risk Rankings'!D10", "Utility Yield"),
        ("Reliance", "RELIANCE.NS", "Conglomerate", "='Risk Rankings'!C12", "='Risk Rankings'!D12", "Index Heavyweight"),
        ("TCS", "TCS.NS", "Information Tech", "='Risk Rankings'!C8", "='Risk Rankings'!D8", "Worst Drawdown"),
        ("NIFTY 50", "^NSEI", "Benchmark Index", "='Risk Rankings'!C16", "='Risk Rankings'!D16", "Benchmark Baseline")
    ]

    for idx, row in enumerate(asset_meta, 9):
        ws1.cell(row=idx, column=1, value=row[0]).font = font_bold if row[0] == "NIFTY 50" else font_regular
        ws1.cell(row=idx, column=2, value=row[1]).font = font_regular
        ws1.cell(row=idx, column=3, value=row[2]).font = font_regular
        
        c_vol = ws1.cell(row=idx, column=4, value=row[3])
        c_vol.font = font_bold if row[0] == "NIFTY 50" else font_regular
        c_vol.number_format = "0.0%"
        c_vol.alignment = align_right
        
        c_mdd = ws1.cell(row=idx, column=5, value=row[4])
        c_mdd.font = font_bold if row[0] == "NIFTY 50" else font_regular
        c_mdd.number_format = "0.0%"
        c_mdd.alignment = align_right
        
        c_cat = ws1.cell(row=idx, column=6, value=row[5])
        c_cat.font = font_regular
        c_cat.alignment = align_center

        for col_idx in range(1, 7):
            ws1.cell(row=idx, column=col_idx).border = border_all
            if idx % 2 == 0:
                ws1.cell(row=idx, column=col_idx).fill = fill_stripe

    auto_fit_columns(ws1, 1, 6)

    # =========================================================================
    # TAB 2: ASSET RISK RANKINGS
    # =========================================================================
    ws2 = wb.create_sheet(title="Risk Rankings")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    ws2["A1"] = "ASSET RISK RANKINGS — VOLATILITY & MAXIMUM DRAWDOWN"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_navy
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells("A2:G2")
    ws2["A2"] = "Quantitative evaluation over 497 trading days (October 2024 to October 2026)"
    ws2["A2"].font = font_subtitle
    ws2["A2"].fill = fill_slate
    ws2["A2"].alignment = align_center

    headers2 = ["Asset Name", "Asset Type", "Annualized Volatility", "Maximum Drawdown", "Peak Date", "Trough Date", "Volatility Rank"]
    for col_idx, h in enumerate(headers2, 1):
        c = ws2.cell(row=5, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = border_all
    ws2.row_dimensions[5].height = 24

    rank_data = [
        ("Infosys", "Equity", 0.282009, -0.4817, "2024-12-13", "2026-07-01"),
        ("M&M", "Equity", 0.280473, -0.2689, "2026-01-02", "2026-10-08"),
        ("TCS", "Equity", 0.251016, -0.5268, "2024-12-13", "2026-07-01"),
        ("L&T", "Equity", 0.246283, -0.2435, "2026-02-23", "2026-03-23"),
        ("NTPC", "Equity", 0.212689, -0.2967, "2024-10-15", "2025-01-13"),
        ("Bharti Airtel", "Equity", 0.208961, -0.1877, "2025-11-21", "2026-05-12"),
        ("Reliance", "Equity", 0.207907, -0.2633, "2026-01-02", "2026-10-01"),
        ("HDFC Bank", "Equity", 0.198052, -0.3104, "2025-07-30", "2026-09-09"),
        ("ITC", "Equity", 0.194978, -0.4237, "2024-10-15", "2026-10-08"),
        ("ICICI Bank", "Equity", 0.184812, -0.1837, "2025-07-23", "2026-03-30"),
        ("NIFTY 50", "Index", 0.131276, -0.1556, "2026-01-02", "2026-10-08")
    ]

    for idx, row in enumerate(rank_data, 6):
        ws2.cell(row=idx, column=1, value=row[0]).font = font_bold if row[1] == "Index" else font_regular
        ws2.cell(row=idx, column=2, value=row[1]).font = font_regular
        ws2.cell(row=idx, column=2).alignment = align_center
        
        c_v = ws2.cell(row=idx, column=3, value=row[2])
        c_v.font = font_bold if row[1] == "Index" else font_regular
        c_v.number_format = "0.0%"
        c_v.alignment = align_right
        
        c_d = ws2.cell(row=idx, column=4, value=row[3])
        c_d.font = font_bold if row[1] == "Index" else font_regular
        c_d.number_format = "0.0%"
        c_d.alignment = align_right
        
        ws2.cell(row=idx, column=5, value=row[4]).alignment = align_center
        ws2.cell(row=idx, column=6, value=row[5]).alignment = align_center
        
        c_r = ws2.cell(row=idx, column=7, value=f"=RANK(C{idx}, $C$6:$C$16)")
        c_r.font = font_bold
        c_r.alignment = align_center

        for col_idx in range(1, 8):
            ws2.cell(row=idx, column=col_idx).border = border_all
            if idx % 2 == 1:
                ws2.cell(row=idx, column=col_idx).fill = fill_stripe

    # Summary Stats Row
    ws2.cell(row=18, column=1, value="Average Equity Volatility").font = font_bold
    ws2.cell(row=18, column=3, value="=AVERAGE(C6:C15)").number_format = "0.0%"
    ws2.cell(row=18, column=3).font = font_bold

    ws2.cell(row=19, column=1, value="Worst Historical Drawdown").font = font_bold
    ws2.cell(row=19, column=4, value="=MIN(D6:D16)").number_format = "0.0%"
    ws2.cell(row=19, column=4).font = font_bold

    auto_fit_columns(ws2, 1, 7)

    # =========================================================================
    # TAB 3: MACRO STRESS TESTING
    # =========================================================================
    ws3 = wb.create_sheet(title="Macro Stress Testing")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:F1")
    ws3["A1"] = "PORTFOLIO MACRO STRESS TESTING SIMULATION"
    ws3["A1"].font = font_title
    ws3["A1"].fill = fill_navy
    ws3["A1"].alignment = align_center
    ws3.row_dimensions[1].height = 28

    ws3.merge_cells("A2:F2")
    ws3["A2"] = "Illustrative ₹100,000 Portfolio across 10 Equal-Weighted Equities (10% each = 100%)"
    ws3["A2"].font = font_subtitle
    ws3["A2"].fill = fill_slate
    ws3["A2"].alignment = align_center

    # Asset Shock Matrix
    headers3 = ["Stock", "Portfolio Weight", "Mild Stress Shock", "Moderate Stress Shock", "Severe Stress Shock"]
    for col_idx, h in enumerate(headers3, 1):
        c = ws3.cell(row=5, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = border_all
    ws3.row_dimensions[5].height = 24

    stress_rows = [
        ("TCS", 0.10, -0.05, -0.12, -0.25),
        ("Infosys", 0.10, -0.05, -0.12, -0.25),
        ("HDFC Bank", 0.10, -0.04, -0.10, -0.22),
        ("ICICI Bank", 0.10, -0.04, -0.10, -0.22),
        ("Reliance", 0.10, -0.05, -0.12, -0.25),
        ("NTPC", 0.10, -0.03, -0.08, -0.18),
        ("L&T", 0.10, -0.05, -0.12, -0.25),
        ("M&M", 0.10, -0.06, -0.15, -0.30),
        ("ITC", 0.10, -0.03, -0.07, -0.15),
        ("Bharti Airtel", 0.10, -0.04, -0.09, -0.20)
    ]

    for idx, row in enumerate(stress_rows, 6):
        ws3.cell(row=idx, column=1, value=row[0]).font = font_bold
        ws3.cell(row=idx, column=2, value=row[1]).number_format = "0.0%"
        ws3.cell(row=idx, column=2).alignment = align_right
        
        for c_idx, val in enumerate(row[2:], 3):
            c_val = ws3.cell(row=idx, column=c_idx, value=val)
            c_val.number_format = "0.0%"
            c_val.alignment = align_right
            c_val.font = font_regular

        for col_idx in range(1, 6):
            ws3.cell(row=idx, column=col_idx).border = border_all
            if idx % 2 == 1:
                ws3.cell(row=idx, column=col_idx).fill = fill_stripe

    # Totals Row
    ws3.cell(row=16, column=1, value="Portfolio Weighted Return").font = font_bold
    ws3.cell(row=16, column=2, value="=SUM(B6:B15)").number_format = "0.0%"
    ws3.cell(row=16, column=2).font = font_bold

    for c_idx, col_letter in enumerate(["C", "D", "E"], 3):
        c_tot = ws3.cell(row=16, column=c_idx, value=f"=SUMPRODUCT($B$6:$B$15, {col_letter}6:{col_letter}15)")
        c_tot.font = font_bold
        c_tot.number_format = "0.0%"
        c_tot.alignment = align_right
        c_tot.border = border_all
        c_tot.fill = fill_subtle_blue

    # Scenario Summary Cards Table
    ws3["A19"] = "Executive Scenario Capital Loss Summary"
    ws3["A19"].font = font_section

    headers_sum = ["Scenario", "Portfolio Return", "Portfolio Loss", "Starting Portfolio", "Post-Stress Residual Value"]
    for col_idx, h in enumerate(headers_sum, 1):
        c = ws3.cell(row=20, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = border_all

    sc_cards = [
        ("Mild Stress", "=C16", "=ABS(B21)", 100000, "=D21*(1+B21)"),
        ("Moderate Stress", "=D16", "=ABS(B22)", 100000, "=D22*(1+B22)"),
        ("Severe Stress", "=E16", "=ABS(B23)", 100000, "=D23*(1+B23)")
    ]

    for idx, row in enumerate(sc_cards, 21):
        ws3.cell(row=idx, column=1, value=row[0]).font = font_bold
        
        c_ret = ws3.cell(row=idx, column=2, value=row[1])
        c_ret.number_format = "0.0%"
        c_ret.font = font_bold
        c_ret.alignment = align_right
        
        c_loss = ws3.cell(row=idx, column=3, value=row[2])
        c_loss.number_format = "0.0%"
        c_loss.font = font_bold
        c_loss.alignment = align_right
        c_loss.fill = fill_red_subtle
        
        c_start = ws3.cell(row=idx, column=4, value=row[3])
        c_start.number_format = "₹#,##0"
        c_start.font = font_regular
        c_start.alignment = align_right
        
        c_post = ws3.cell(row=idx, column=5, value=row[4])
        c_post.number_format = "₹#,##0"
        c_post.font = font_bold
        c_post.alignment = align_right

        for col_idx in range(1, 6):
            ws3.cell(row=idx, column=col_idx).border = border_all

    auto_fit_columns(ws3, 1, 5)

    # =========================================================================
    # TAB 4: OPTIONS INTELLIGENCE
    # =========================================================================
    ws4 = wb.create_sheet(title="Options Intelligence")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:G1")
    ws4["A1"] = "OPTIONS STRATEGY INTELLIGENCE & EXPIRATION PAYOFF"
    ws4["A1"].font = font_title
    ws4["A1"].fill = fill_navy
    ws4["A1"].alignment = align_center
    ws4.row_dimensions[1].height = 28

    ws4.merge_cells("A2:G2")
    ws4["A2"] = "Payoff curves across tested underlying strike ranges (₹700 to ₹1,300)"
    ws4["A2"].font = font_subtitle
    ws4["A2"].fill = fill_slate
    ws4["A2"].alignment = align_center

    # Strategy Overview
    ws4["A4"] = "Strategy Specifications & Risk Profile"
    ws4["A4"].font = font_section

    headers4_spec = ["Strategy", "Description", "Net Premium Paid", "Max Profit", "Max Loss", "Lower Breakeven", "Upper Breakeven"]
    for col_idx, h in enumerate(headers4_spec, 1):
        c = ws4.cell(row=5, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = border_all

    opt_specs = [
        ("Bull Call Spread", "Limited-risk bullish play (Long 1000C / Short 1100C)", 35.0, 65.0, -35.0, 1035.0, 1035.0),
        ("Long Straddle", "Volatility breakout play (Long 1000C / Long 1000P)", 105.0, "Substantial", -105.0, 895.0, 1105.0),
        ("Iron Condor", "Range-bound credit play (850P/900P spread + 1100C/1150C spread)", -23.0, 23.0, -27.0, 877.0, 1123.0)
    ]

    for idx, row in enumerate(opt_specs, 6):
        ws4.cell(row=idx, column=1, value=row[0]).font = font_bold
        ws4.cell(row=idx, column=2, value=row[1]).font = font_regular
        
        c_p = ws4.cell(row=idx, column=3, value=row[2])
        c_p.number_format = "₹#,##0.0"
        c_p.alignment = align_right
        
        c_mp = ws4.cell(row=idx, column=4, value=row[3])
        if isinstance(row[3], (int, float)):
            c_mp.number_format = "₹#,##0.0"
        c_mp.alignment = align_right
        c_mp.font = font_bold
        
        c_ml = ws4.cell(row=idx, column=5, value=row[4])
        c_ml.number_format = "₹#,##0.0"
        c_ml.alignment = align_right
        c_ml.font = font_bold
        
        c_bl = ws4.cell(row=idx, column=6, value=row[5])
        c_bl.number_format = "₹#,##0.0"
        c_bl.alignment = align_right
        
        c_bu = ws4.cell(row=idx, column=7, value=row[6])
        c_bu.number_format = "₹#,##0.0"
        c_bu.alignment = align_right

        for col_idx in range(1, 8):
            ws4.cell(row=idx, column=col_idx).border = border_all

    # Expiration Payoff Schedule Table
    ws4["A11"] = "Expiration Profit / Loss Schedule (₹ per unit)"
    ws4["A11"].font = font_section

    headers4_pay = ["Spot Price at Expiry", "Bull Call Spread", "Long Straddle", "Iron Condor"]
    for col_idx, h in enumerate(headers4_pay, 1):
        c = ws4.cell(row=12, column=col_idx, value=h)
        c.font = font_header
        c.fill = fill_navy
        c.alignment = align_center
        c.border = border_all

    # Spot prices 700 to 1300 at step 25
    prices = list(range(700, 1325, 25))
    for r_idx, s in enumerate(prices, 13):
        # Formulas:
        # BCS: =MAX(A{r_idx}-1000, 0) - MAX(A{r_idx}-1100, 0) - 35
        # Straddle: =MAX(A{r_idx}-1000, 0) + MAX(1000-A{r_idx}, 0) - 105
        # IC: =MAX(850-A{r_idx}, 0) - MAX(900-A{r_idx}, 0) - MAX(A{r_idx}-1100, 0) + MAX(A{r_idx}-1150, 0) + 23
        ws4.cell(row=r_idx, column=1, value=s).number_format = "₹#,##0"
        ws4.cell(row=r_idx, column=1).alignment = align_right
        ws4.cell(row=r_idx, column=1).font = font_bold

        f_bcs = f"=MAX(A{r_idx}-1000, 0) - MAX(A{r_idx}-1100, 0) - 35"
        c_bcs = ws4.cell(row=r_idx, column=2, value=f_bcs)
        c_bcs.number_format = "₹#,##0"
        c_bcs.alignment = align_right

        f_ls = f"=MAX(A{r_idx}-1000, 0) + MAX(1000-A{r_idx}, 0) - 105"
        c_ls = ws4.cell(row=r_idx, column=3, value=f_ls)
        c_ls.number_format = "₹#,##0"
        c_ls.alignment = align_right

        f_ic = f"=MAX(850-A{r_idx}, 0) - MAX(900-A{r_idx}, 0) - MAX(A{r_idx}-1100, 0) + MAX(A{r_idx}-1150, 0) + 23"
        c_ic = ws4.cell(row=r_idx, column=4, value=f_ic)
        c_ic.number_format = "₹#,##0"
        c_ic.alignment = align_right

        for col_idx in range(1, 5):
            ws4.cell(row=r_idx, column=col_idx).border = border_all
            if r_idx % 2 == 1:
                ws4.cell(row=r_idx, column=col_idx).fill = fill_stripe

    auto_fit_columns(ws4, 1, 7)

    # =========================================================================
    # TAB 5: CORRELATION MATRIX
    # =========================================================================
    ws5 = wb.create_sheet(title="Correlation Matrix")
    ws5.views.sheetView[0].showGridLines = True

    ws5.merge_cells("A1:L1")
    ws5["A1"] = "PAIRWISE CORRELATION HEATMAP (DAILY RETURNS)"
    ws5["A1"].font = font_title
    ws5["A1"].fill = fill_navy
    ws5["A1"].alignment = align_center
    ws5.row_dimensions[1].height = 28

    ws5.merge_cells("A2:L2")
    ws5["A2"] = "Pearson correlation coefficients across 495 return observations (Aligned dates)"
    ws5["A2"].font = font_subtitle
    ws5["A2"].fill = fill_slate
    ws5["A2"].alignment = align_center

    assets_11 = [
        "Bharti Airtel", "HDFC Bank", "ICICI Bank", "Infosys", "ITC",
        "L&T", "M&M", "NTPC", "Reliance", "TCS", "NIFTY 50"
    ]

    # Column Headers
    ws5.cell(row=4, column=1, value="Asset").font = font_header
    ws5.cell(row=4, column=1).fill = fill_navy
    ws5.cell(row=4, column=1).border = border_all
    ws5.cell(row=4, column=1).alignment = align_center

    for c_idx, a_name in enumerate(assets_11, 2):
        c = ws5.cell(row=4, column=c_idx, value=a_name)
        c.font = font_header
        c.fill = fill_navy
        c.border = border_all
        c.alignment = align_center
    ws5.row_dimensions[4].height = 24

    # Correlation Matrix Values
    # We load symmetric matrix values
    corr_data = {
        "Bharti Airtel": {"Bharti Airtel": 1.0, "HDFC Bank": 0.2647, "ICICI Bank": 0.2631, "Infosys": 0.1781, "ITC": 0.2716, "L&T": 0.2451, "M&M": 0.3198, "NTPC": 0.1857, "Reliance": 0.3188, "TCS": 0.1601, "NIFTY 50": 0.4881},
        "HDFC Bank": {"Bharti Airtel": 0.2647, "HDFC Bank": 1.0, "ICICI Bank": 0.5038, "Infosys": 0.2057, "ITC": 0.2534, "L&T": 0.4593, "M&M": 0.4084, "NTPC": 0.2325, "Reliance": 0.3887, "TCS": 0.1933, "NIFTY 50": 0.7091},
        "ICICI Bank": {"Bharti Airtel": 0.2631, "HDFC Bank": 0.5038, "ICICI Bank": 1.0, "Infosys": 0.1282, "ITC": 0.2119, "L&T": 0.4008, "M&M": 0.4117, "NTPC": 0.273, "Reliance": 0.2773, "TCS": 0.172, "NIFTY 50": 0.6356},
        "Infosys": {"Bharti Airtel": 0.1781, "HDFC Bank": 0.2057, "ICICI Bank": 0.1282, "Infosys": 1.0, "ITC": 0.1093, "L&T": 0.2156, "M&M": 0.2288, "NTPC": 0.0465, "Reliance": 0.1923, "TCS": 0.7448, "NIFTY 50": 0.417},
        "ITC": {"Bharti Airtel": 0.2716, "HDFC Bank": 0.2534, "ICICI Bank": 0.2119, "Infosys": 0.1093, "ITC": 1.0, "L&T": 0.2657, "M&M": 0.2552, "NTPC": 0.1798, "Reliance": 0.3222, "TCS": 0.1962, "NIFTY 50": 0.4426},
        "L&T": {"Bharti Airtel": 0.2451, "HDFC Bank": 0.4593, "ICICI Bank": 0.4008, "Infosys": 0.2156, "ITC": 0.2657, "L&T": 1.0, "M&M": 0.3932, "NTPC": 0.4249, "Reliance": 0.4431, "TCS": 0.2582, "NIFTY 50": 0.6988},
        "M&M": {"Bharti Airtel": 0.3198, "HDFC Bank": 0.4084, "ICICI Bank": 0.4117, "Infosys": 0.2288, "ITC": 0.2552, "L&T": 0.3932, "M&M": 1.0, "NTPC": 0.2577, "Reliance": 0.3577, "TCS": 0.2333, "NIFTY 50": 0.6492},
        "NTPC": {"Bharti Airtel": 0.1857, "HDFC Bank": 0.2325, "ICICI Bank": 0.273, "Infosys": 0.0465, "ITC": 0.1798, "L&T": 0.4249, "M&M": 0.2577, "NTPC": 1.0, "Reliance": 0.3715, "TCS": 0.065, "NIFTY 50": 0.4912},
        "Reliance": {"Bharti Airtel": 0.3188, "HDFC Bank": 0.3887, "ICICI Bank": 0.2773, "Infosys": 0.1923, "ITC": 0.3222, "L&T": 0.4431, "M&M": 0.3577, "NTPC": 0.3715, "Reliance": 1.0, "TCS": 0.2427, "NIFTY 50": 0.6567},
        "TCS": {"Bharti Airtel": 0.1601, "HDFC Bank": 0.1933, "ICICI Bank": 0.172, "Infosys": 0.7448, "ITC": 0.1962, "L&T": 0.2582, "M&M": 0.2333, "NTPC": 0.065, "Reliance": 0.2427, "TCS": 1.0, "NIFTY 50": 0.4399},
        "NIFTY 50": {"Bharti Airtel": 0.4881, "HDFC Bank": 0.7091, "ICICI Bank": 0.6356, "Infosys": 0.417, "ITC": 0.4426, "L&T": 0.6988, "M&M": 0.6492, "NTPC": 0.4912, "Reliance": 0.6567, "TCS": 0.4399, "NIFTY 50": 1.0}
    }

    for r_idx, r_name in enumerate(assets_11, 5):
        # Row header
        ws5.cell(row=r_idx, column=1, value=r_name).font = font_bold
        ws5.cell(row=r_idx, column=1).border = border_all
        ws5.cell(row=r_idx, column=1).fill = fill_subtle_blue

        for c_idx, c_name in enumerate(assets_11, 2):
            val = corr_data[r_name][c_name]
            cell = ws5.cell(row=r_idx, column=c_idx, value=val)
            cell.number_format = "0.00"
            cell.alignment = align_center
            cell.border = border_all
            
            # Subtle Heatmap Fill
            if r_name == c_name:
                cell.fill = PatternFill(start_color="CBD5E0", end_color="CBD5E0", fill_type="solid")
                cell.font = font_bold
            elif val >= 0.60:
                cell.fill = PatternFill(start_color="FED7D7", end_color="FED7D7", fill_type="solid")  # High co-movement (red)
                cell.font = font_bold
            elif val <= 0.15:
                cell.fill = PatternFill(start_color="C6F6D5", end_color="C6F6D5", fill_type="solid")  # High diversification (green)
                cell.font = font_bold
            else:
                cell.font = font_regular

    auto_fit_columns(ws5, 1, 12)

    # Save
    wb.save(EXCEL_PATH)
    print(f"Financial Risk Model Workbook successfully built at: {EXCEL_PATH}")

if __name__ == "__main__":
    create_financial_workbook()
